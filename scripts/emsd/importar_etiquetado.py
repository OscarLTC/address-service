# Convierte la planilla de etiquetado revisada al formato del dataset de oro, con la
# columna opcional expected_interior. Lee y escribe solo dentro de data/private/:
# contiene direcciones de clientes.
#
#   python scripts/emsd/importar_etiquetado.py \
#       data/private/etiquetado_muestra_v1_completado.xlsx data/private/golden_real_v1.csv \
#       --metodo llm_chatgpt --revisor "ChatGPT (sin revisión humana)"
#
# Solo entran las filas CORRECTO y CORREGIDO. DUDOSO y NO ES DIRECCIÓN se cuentan pero
# no se evalúan: no tienen una verdad definida.
import argparse, csv, os, re, sys
from collections import Counter
from openpyxl import load_workbook

BASE = ["id", "source", "raw_address", "district_field", "province_field", "department_field",
        "expected_street_type", "expected_street_name", "expected_number", "expected_block",
        "expected_lot", "expected_urbanization", "expected_ubigeo", "lat_verified", "lng_verified",
        "verification_method", "verified_by", "split", "notes", "expected_interior"]

# Columnas de la hoja Etiquetado (fila 2 = encabezados; ver planilla_etiquetado.py).
COL = {"id": 1, "direccion": 3, "numero": 4, "distrito": 6, "provincia": 7,
       "p_numero": 11,
       "tipo": 19, "nombre": 20, "num": 21, "interior": 22, "mz": 23, "lt": 24, "urb": 25, "ubigeo": 26,
       "estado": 27, "comentario": 29, "particion": 31}
RE_SN = re.compile(r"\bS\s*/\s*N\b", re.I)


def private(path):
    return os.path.normpath(path).replace("\\", "/").startswith("data/private/")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("out")
    ap.add_argument("--metodo", required=True, help="verification_method (p. ej. humano, llm_chatgpt)")
    ap.add_argument("--revisor", required=True, help="verified_by")
    a = ap.parse_args()
    if not (private(a.xlsx) and private(a.out)):
        sys.exit("la planilla y la salida deben estar dentro de data/private/")

    ws = load_workbook(a.xlsx)["Etiquetado"]
    head = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    if head[COL["estado"] - 1] != "Estado" or head[COL["direccion"] - 1] != "Dirección (EMS-D)":
        sys.exit("la hoja Etiquetado no tiene el formato esperado")

    def v(r, k):
        x = ws.cell(row=r, column=COL[k]).value
        return "" if x is None else str(x).strip()

    estados, ajustes, out = Counter(), Counter(), []
    for r in range(3, ws.max_row + 1):
        if not v(r, "id"):
            continue
        estado = v(r, "estado")
        estados[estado] += 1
        if estado not in ("CORRECTO", "CORREGIDO"):
            continue
        raw = v(r, "direccion") + (" " + v(r, "numero") if v(r, "numero") else "")
        number, notes = v(r, "num"), [f"estado={estado}"]
        # La instrucción pedía "solo dígitos" y omitió S/N: si el texto dice S/N y la
        # propuesta lo tenía, se restaura.
        if not number and v(r, "p_numero") == "S/N" and RE_SN.search(raw):
            number = "S/N"
            notes.append("S/N restaurado")
            ajustes["S/N restaurado"] += 1
        if v(r, "comentario"):
            notes.append(v(r, "comentario"))
        out.append({
            "id": "EMS-" + v(r, "id"), "source": "emsd_real", "raw_address": raw,
            "district_field": v(r, "distrito"), "province_field": v(r, "provincia"), "department_field": "",
            "expected_street_type": v(r, "tipo"), "expected_street_name": v(r, "nombre"),
            "expected_number": number, "expected_block": v(r, "mz"), "expected_lot": v(r, "lt"),
            "expected_urbanization": v(r, "urb"), "expected_ubigeo": v(r, "ubigeo"),
            "lat_verified": "", "lng_verified": "", "verification_method": a.metodo,
            "verified_by": a.revisor, "split": v(r, "particion"), "notes": "; ".join(notes),
            "expected_interior": v(r, "interior"),
        })

    with open(a.out, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=BASE)
        w.writeheader()
        w.writerows(out)
    print(f"{a.out}: {len(out)} filas | estados: {dict(estados)} | ajustes: {dict(ajustes)}")


if __name__ == "__main__":
    main()
